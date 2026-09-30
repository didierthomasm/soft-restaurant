"""Attendance workbook: calendar grid, HR incident list and per-week summary."""

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from tabernas.domain.periods import days
from tabernas.domain.summary import EmployeeSummary
from tabernas.domain.types import DayResult
from tabernas.export.labels import DAY_ABBR, OUTCOME_FILLS, OUTCOME_LABELS, RH_LABELS
from tabernas.services.attendance import AttendanceReport

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SUMMARY_HEADER = (
    "Empleado",
    "Periodo",
    "Días trabajados",
    "Retardos",
    "Retardos justificados",
    "Faltas",
    "Faltas justificadas",
    "Pendientes de resolver",
)


def cell_text(result: DayResult) -> str:
    label = OUTCOME_LABELS[result.outcome]
    if result.checkin is None:
        return label
    return f"{label} {result.checkin:%H:%M}".strip()


def build_workbook(report: AttendanceReport, summaries: Sequence[EmployeeSummary]) -> bytes:
    workbook = Workbook()
    calendar = workbook.active
    assert calendar is not None
    _calendar_sheet(calendar, report)
    _rh_sheet(workbook.create_sheet("Incidencias RH"), report)
    _summary_sheet(workbook.create_sheet("Resumen"), report, summaries)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _header(sheet: Worksheet, titles: Sequence[str]) -> None:
    sheet.append(list(titles))
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def _calendar_sheet(sheet: Worksheet, report: AttendanceReport) -> None:
    sheet.title = "Calendario"
    period = days(report.start, report.end)
    _header(sheet, ["Empleado", *(f"{DAY_ABBR[d.weekday()]} {d:%d/%m}" for d in period)])
    by_key = {(r.employee_id, r.day): r for r in report.results}
    for employee in report.employees:
        results = [by_key[(employee.id, d)] for d in period]
        sheet.append([employee.short_name, *(cell_text(r) for r in results)])
        for column, result in enumerate(results, start=2):
            sheet.cell(row=sheet.max_row, column=column).fill = PatternFill(
                "solid", fgColor=OUTCOME_FILLS[result.outcome]
            )


def _rh_sheet(sheet: Worksheet, report: AttendanceReport) -> None:
    _header(sheet, ["Nombre en RH", "Fecha", "Tipo", "Comentario"])
    for row in report.rh_rows:
        sheet.append([row.name, row.day, RH_LABELS[row.rh_type], row.comment])
        sheet.cell(row=sheet.max_row, column=2).number_format = "dd/mm/yyyy"


def _summary_sheet(
    sheet: Worksheet, report: AttendanceReport, summaries: Sequence[EmployeeSummary]
) -> None:
    _header(sheet, SUMMARY_HEADER)
    names = {e.id: e.short_name for e in report.employees}
    for s in summaries:
        sheet.append(
            [
                names.get(s.employee_id, str(s.employee_id)),
                s.period,
                s.worked,
                s.late,
                s.late_justified,
                s.absent,
                s.absent_justified,
                s.unresolved,
            ]
        )
