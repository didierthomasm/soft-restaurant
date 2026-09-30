"""Reads SR's attendance export for reconciliation. SR's '.XLS' files are really xlsx."""

from collections import Counter
from collections.abc import Iterator, Sequence
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

CLAVE_HEADER = "CLAVEEMPLEADO"
ENTRADA_HEADER = "ENTRADA"


def read_attendance_export(path: Path, start: date, end: date) -> Counter[int]:
    # Load from bytes: openpyxl refuses the .XLS extension even though the content is xlsx.
    workbook = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    clave_col, entrada_col = _find_header(rows)
    counts: Counter[int] = Counter()
    for row in rows:  # the iterator continues after the header row
        if len(row) <= max(clave_col, entrada_col):
            continue
        clave, entrada = row[clave_col], row[entrada_col]
        in_range = isinstance(entrada, datetime) and start <= entrada.date() <= end
        if isinstance(clave, int) and in_range:
            counts[clave] += 1
    return counts


def _find_header(rows: Iterator[Sequence[object]]) -> tuple[int, int]:
    for row in rows:
        values = list(row)
        if CLAVE_HEADER in values and ENTRADA_HEADER in values:
            return values.index(CLAVE_HEADER), values.index(ENTRADA_HEADER)
    raise ValueError("El archivo no parece un export de asistencia de SR")
