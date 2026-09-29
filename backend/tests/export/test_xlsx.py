from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from tabernas.domain.summary import Grouping, summarize
from tabernas.domain.types import DayResult, Outcome, Planned, RhRow, RhType
from tabernas.export.labels import OUTCOME_FILLS, OUTCOME_LABELS, RH_LABELS
from tabernas.export.xlsx import build_workbook
from tabernas.services.attendance import AttendanceReport
from tests.domain.factories import employee


def sample_report() -> AttendanceReport:
    results = (
        DayResult(
            1, date(2026, 9, 21), Planned.WORK, Outcome.OK, checkin=datetime(2026, 9, 21, 16, 35)
        ),
        DayResult(1, date(2026, 9, 22), Planned.REST, Outcome.REST),
        DayResult(
            1,
            date(2026, 9, 23),
            Planned.WORK,
            Outcome.LATE,
            checkin=datetime(2026, 9, 23, 16, 55),
            minutes_late=15,
        ),
    )
    return AttendanceReport(
        start=date(2026, 9, 21),
        end=date(2026, 9, 23),
        employees=(employee(1, rh_name="APELLIDO UNO"),),
        results=results,
        rh_rows=(RhRow(1, "APELLIDO UNO", date(2026, 9, 23), RhType.RETARDO, ""),),
        warnings=(),
    )


def values(row: tuple) -> list[object]:
    return [cell.value for cell in row]


def test_workbook_sheets_and_content() -> None:
    report = sample_report()
    content = build_workbook(report, summarize(report.results, Grouping.WEEK))
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Calendario", "Incidencias RH", "Resumen"]

    calendar = workbook["Calendario"]
    assert values(calendar[1]) == ["Empleado", "Lun 21/09", "Mar 22/09", "Mié 23/09"]
    assert values(calendar[2]) == ["E1", "A tiempo 16:35", "Descanso", "Retardo 16:55"]

    rh = workbook["Incidencias RH"]
    assert values(rh[1]) == ["Nombre en RH", "Fecha", "Tipo", "Comentario"]
    assert values(rh[2])[:3] == ["APELLIDO UNO", datetime(2026, 9, 23), "Retardo"]

    summary = workbook["Resumen"]
    assert values(summary[1]) == [
        "Empleado",
        "Periodo",
        "Días trabajados",
        "Retardos",
        "Retardos justificados",
        "Faltas",
        "Faltas justificadas",
        "Pendientes de resolver",
    ]
    assert values(summary[2]) == ["E1", "2026-W39", 2, 1, 0, 0, 0, 0]


def test_labels_cover_every_enum_value() -> None:
    assert set(OUTCOME_LABELS) == set(Outcome)
    assert set(OUTCOME_FILLS) == set(Outcome)
    assert set(RH_LABELS) == set(RhType)
