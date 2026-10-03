from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from tabernas.demo import seed_demo_data
from tabernas.sr.fake_source import FakeSource
from tabernas.sr.source import SrUnavailableError
from tests.api.conftest import MakeClient
from tests.support import StubSource

WEEK = {"from": "2026-09-21", "to": "2026-09-27"}
INCIDENT_OUTCOMES = {"LATE", "ABSENT", "UNREGISTERED_CHANGE", "JUSTIFIED"}


@pytest.fixture
def demo_client(session_factory: sessionmaker[Session], make_client: MakeClient) -> TestClient:
    with session_factory() as session:
        seed_demo_data(session)
        session.commit()
    return make_client(FakeSource(today=lambda: date(2026, 9, 27)))


def test_calendar_has_one_day_per_employee(demo_client: TestClient) -> None:
    data = demo_client.get("/attendance/calendar", params=WEEK).json()["data"]
    assert (data["start"], data["end"]) == ("2026-09-21", "2026-09-27")
    assert len(data["employees"]) == 7
    assert len(data["days"]) == 49
    assert not {w["code"] for w in data["warnings"]} & {"NO_REST_RULE", "UNMAPPED_CHECKIN"}


def test_calendar_days_carry_their_exception(demo_client: TestClient) -> None:
    employee_id = demo_client.get("/employees").json()["data"][0]["id"]
    body = {
        "kind": "WORK_TO_ABSENCE",
        "employee_id": employee_id,
        "date_from": "2026-09-23",
        "date_to": "2026-09-24",
        "rh_type": "VACACIONES",
        "comment": "Viaje",
    }
    created = demo_client.post("/exceptions", json=body).json()["data"]
    days = demo_client.get("/attendance/calendar", params=WEEK).json()["data"]["days"]
    own = {d["day"]: d for d in days if d["employee_id"] == employee_id}
    expected = {k: created[k] for k in ("id", "kind", "date_from", "date_to", "rh_type", "comment")}
    assert own["2026-09-23"]["exception"] == expected
    assert own["2026-09-24"]["exception"] == expected
    assert own["2026-09-21"]["exception"] is None


def test_incidents_and_rh_rows_are_consistent(demo_client: TestClient) -> None:
    data = demo_client.get("/attendance/incidents", params=WEEK).json()["data"]
    assert {d["outcome"] for d in data["incidents"]} <= INCIDENT_OUTCOMES
    incident_keys = {(d["employee_id"], d["day"]) for d in data["incidents"]}
    assert data["rh_rows"]
    assert {(r["employee_id"], r["day"]) for r in data["rh_rows"]} <= incident_keys


def test_monthly_summary(demo_client: TestClient) -> None:
    params = {"from": "2026-09-01", "to": "2026-09-30", "group": "month"}
    data = demo_client.get("/attendance/summary", params=params).json()["data"]
    assert [s["period"] for s in data] == ["2026-09"] * 7
    assert all(isinstance(s["justified_by_type"], dict) for s in data)


def test_bad_ranges_are_422(client: TestClient) -> None:
    assert client.get("/attendance/calendar", params={"from": "2026-09-21"}).status_code == 422
    reversed_range = {"from": "2026-09-27", "to": "2026-09-21"}
    assert client.get("/attendance/calendar", params=reversed_range).status_code == 422
    too_long = {"from": "2026-01-01", "to": "2026-12-31"}
    assert client.get("/attendance/incidents", params=too_long).status_code == 422


def test_sr_down_is_503(make_client: MakeClient) -> None:
    client = make_client(StubSource(error=SrUnavailableError("x")))
    response = client.get("/attendance/calendar", params=WEEK)
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "SR_UNAVAILABLE"


def test_export_xlsx(demo_client: TestClient) -> None:
    response = demo_client.get("/attendance/export.xlsx", params=WEEK)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers["content-disposition"]
    assert 'filename="asistencia_2026-09-21_2026-09-27.xlsx"' in disposition
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["Calendario"].max_row == 8  # header + 7 employees


def test_export_xlsx_summary_grouped_by_month(demo_client: TestClient) -> None:
    params = {"from": "2026-09-01", "to": "2026-09-30", "group": "month"}
    response = demo_client.get("/attendance/export.xlsx", params=params)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    resumen = workbook["Resumen"]
    periods = {str(row[1].value) for row in resumen.iter_rows(min_row=2)}
    assert periods == {"2026-09"}
